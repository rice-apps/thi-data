"""
File loader strategy/registry for the ETL pipeline.

Each loader creates a DuckDB table with all-VARCHAR columns from a given
file path, matching the behavior of DuckDB's read_csv(..., all_varchar=True).

Adding a new format: one function + one @register_loader(".ext") decorator.
"""

import logging
from datetime import date, datetime
from pathlib import Path
from typing import Callable

import duckdb
import pyarrow as pa

logger = logging.getLogger(__name__)

_LOADER_REGISTRY: dict[str, Callable[[duckdb.DuckDBPyConnection, str, str], None]] = {}


def register_loader(extension: str):
    """Decorator to register a loader function for a file extension."""
    def decorator(fn):
        _LOADER_REGISTRY[extension.lower()] = fn
        return fn
    return decorator


def create_raw_table(
    con: duckdb.DuckDBPyConnection,
    file_path: str,
    table_name: str,
) -> None:
    ext = Path(file_path).suffix.lower()
    loader = _LOADER_REGISTRY.get(ext)
    if loader is None:
        supported = ", ".join(sorted(_LOADER_REGISTRY.keys()))
        raise ValueError(
            f"Unsupported file format '{ext}'. Supported formats: {supported}"
        )
    logger.info("Loading '%s' file via %s: %s", ext, loader.__name__, file_path)
    loader(con, file_path, table_name)


@register_loader(".csv")
def _load_csv(
    con: duckdb.DuckDBPyConnection,
    file_path: str,
    table_name: str,
) -> None:
    con.execute(f"""
        CREATE TABLE {table_name} AS
        SELECT * FROM read_csv('{file_path}', all_varchar=True, auto_detect=True)
    """)


def _cell_to_str(val):
    """Convert an openpyxl cell value to a string, handling dates and datetimes."""
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.isoformat()
    if isinstance(val, date):
        return val.strftime("%Y-%m-%d")
    return str(val)


@register_loader(".xlsx")
def _load_xlsx(
    con: duckdb.DuckDBPyConnection,
    file_path: str,
    table_name: str,
) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(file_path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)

        header_row = next(rows_iter, None)
        if header_row is None:
            raise ValueError(f"XLSX file is empty or has no header: {file_path}")

        headers = []
        for i, val in enumerate(header_row):
            if val is None:
                raise ValueError(
                    f"XLSX header contains empty cell in column {i + 1}. "
                    "All header cells must have values."
                )
            headers.append(str(val).strip())

        columns: list[list[str | None]] = [[] for _ in headers]
        for row in rows_iter:
            for i, val in enumerate(row):
                if i < len(headers):
                    columns[i].append(_cell_to_str(val))
    finally:
        wb.close()

    arrow_table = pa.table(
        {h: pa.array(col, type=pa.string()) for h, col in zip(headers, columns)}
    )

    con.register("_arrow_staging", arrow_table)
    try:
        con.execute(f"CREATE TABLE {table_name} AS SELECT * FROM _arrow_staging")
    finally:
        con.unregister("_arrow_staging")
